"""
쇼핑검색광고 리포트(NBOS 매칭, 상품 단위 원본) -> 카테고리/브랜드별 일별 집계 변환기.

원본은 상품 x 디바이스 x 광고그룹 단위로 매우 세분화되어 있고(하루 5만 행 이상,
파일 하나가 수백MB) 여러 달치를 한 파일에 월별 시트로 나눠 담는 경우도 있어(예:
"25년_1월", "25년_2월", ...), pandas.read_excel로 통째로 읽으면 느리고 메모리를
많이 먹는다. openpyxl의 read_only 스트리밍 모드로 행 단위로 훑으면서 바로
집계(dict 누적)하기 때문에 대용량 파일도 안전하게 처리한다.

날짜 x 대카테고리 x 중카테고리 x 브랜드명 x 자사/입점 단위로 노출수/클릭수/광고비/
구매수량/판매액을 합산해 가벼운 CSV로 만든다. 원본 컬럼 구성은 시기별로 조금씩
다르지만(2025년 파일은 27개 컬럼, 2026년 7월 이후는 37개 컬럼 — 전환수·전환매출액 등
추가), 우리가 쓰는 컬럼(날짜/대카테고리/중카테고리/브랜드명/자사입점/노출수/클릭수/
총비용/구매수량/판매액)은 전 기간 공통으로 존재해서 컬럼명 기준으로 안전하게 뽑는다.

주의: 원본의 "전환매출액(원)"(간접전환 포함, 광고 플랫폼 자체 귀속 기준)은 쓰지 않는다.
실제 매출과 크게 괴리될 수 있어(관찰된 사례: 판매액의 10배 이상), 효율(ROAS) 계산에는
반드시 "판매액"(NBOS 매칭 실매출)만 사용한다. 애초에 이 컬럼을 집계 대상에서 뺐다.

사용법:
    python build_ad_product_daily.py <원본1.xlsx> [<원본2.xlsx> ...] [--out ad_product_category_daily.csv]

여러 파일을 한 번에 넘기면 전부 훑어서 하나로 합쳐 집계한다(파일마다 시트가
여러 개여도 전부 처리). 이미 출력 파일이 있으면, 새로 변환한 날짜 구간은 새
값으로 덮어쓰고(재확정 반영), 그 외 기존 날짜는 그대로 유지한 채 합쳐서 저장한다
— 기간을 나눠 여러 번 실행해도 안전하다.
"""

import argparse
import os
import sys
from collections import defaultdict
from datetime import datetime, date

import openpyxl
import pandas as pd

REQUIRED_COLS = [
    "날짜", "대카테고리", "중카테고리", "브랜드명", "자사/입점",
    "노출수", "클릭수", "총비용(VAT포함,원)", "구매수량", "판매액",
]
NUMERIC_KEYS = ["노출수", "클릭수", "총비용(VAT포함,원)", "구매수량", "판매액"]


def _to_num(v):
    if v is None or v == "-":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _process_sheet(ws, acc, sheet_label):
    col_idx = {}
    n_rows = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            for j, name in enumerate(row):
                if name in REQUIRED_COLS:
                    col_idx[name] = j
            missing = [c for c in REQUIRED_COLS if c not in col_idx]
            if missing:
                raise ValueError(f"[{sheet_label}] 컬럼을 찾을 수 없습니다: {missing}")
            continue

        date_val = row[col_idx["날짜"]]
        if date_val is None or date_val == "-":
            continue
        if isinstance(date_val, datetime):
            date_key = date_val.date()
        elif isinstance(date_val, date):
            date_key = date_val
        else:
            try:
                date_key = pd.to_datetime(date_val).date()
            except Exception:
                continue

        key = (
            date_key,
            row[col_idx["대카테고리"]],
            row[col_idx["중카테고리"]],
            row[col_idx["브랜드명"]],
            row[col_idx["자사/입점"]],
        )
        rec = acc[key]
        rec[0] += _to_num(row[col_idx["노출수"]])
        rec[1] += _to_num(row[col_idx["클릭수"]])
        rec[2] += _to_num(row[col_idx["총비용(VAT포함,원)"]])
        rec[3] += _to_num(row[col_idx["구매수량"]])
        rec[4] += _to_num(row[col_idx["판매액"]])
        n_rows += 1

    return n_rows


def convert(raw_paths: list, out_path: str):
    acc = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0, 0.0])  # 노출수,클릭수,광고비,구매수량,판매액
    total_rows = 0

    for raw_path in raw_paths:
        print(f"원본 로드 중: {raw_path}")
        wb = openpyxl.load_workbook(raw_path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            n = _process_sheet(ws, acc, f"{os.path.basename(raw_path)}::{sheet_name}")
            total_rows += n
            print(f"  시트 '{sheet_name}': {n:,}행 처리 (누적 {total_rows:,}행)")
        wb.close()

    if not acc:
        print("[오류] 집계된 데이터가 없습니다.")
        sys.exit(1)

    rows = []
    for (date_key, large_cat, mid_cat, brand, own), (imp, clk, cost, qty, sale) in acc.items():
        rows.append({
            "date": date_key, "대카테고리": large_cat, "중카테고리": mid_cat,
            "브랜드명": brand, "자사/입점": own,
            "노출수": imp, "클릭수": clk, "광고비": cost, "구매수량": qty, "판매액": sale,
        })
    agg = pd.DataFrame(rows)
    agg["date"] = pd.to_datetime(agg["date"])
    print(f"집계 완료: {total_rows:,}행 -> {len(agg):,}행 "
          f"({agg['date'].min().date()} ~ {agg['date'].max().date()}, "
          f"대카테고리 {agg['대카테고리'].nunique()}개, 브랜드 {agg['브랜드명'].nunique()}개)")

    if os.path.exists(out_path):
        existing = pd.read_csv(out_path, parse_dates=["date"], encoding="utf-8-sig")
        new_dates = set(agg["date"].unique())
        kept = existing[~existing["date"].isin(new_dates)]
        merged = pd.concat([kept, agg], ignore_index=True)
        print(f"기존 파일 병합: 기존 {len(existing):,}행 중 겹치는 날짜 제외 {len(kept):,}행 유지 "
              f"+ 신규 {len(agg):,}행 -> 총 {len(merged):,}행")
    else:
        merged = agg

    merged = merged.sort_values(["date", "대카테고리", "중카테고리", "브랜드명"]).reset_index(drop=True)
    merged.to_csv(out_path, index=False, encoding="utf-8-sig")
    print(f"저장 완료: {out_path} "
          f"(전체 날짜 범위: {merged['date'].min().date()} ~ {merged['date'].max().date()})")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("raw_paths", nargs="+", help="원본 엑셀 파일 경로(여러 개 가능, 각 파일의 모든 시트를 처리)")
    parser.add_argument("--out", default="ad_product_category_daily.csv", help="출력 CSV 경로")
    args = parser.parse_args()
    convert(args.raw_paths, args.out)
